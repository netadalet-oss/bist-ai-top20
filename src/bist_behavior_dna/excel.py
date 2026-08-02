from __future__ import annotations

from pathlib import Path
from typing import Mapping

import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """Export analysis outputs to an audit-friendly multi-sheet workbook."""

    def export(self, path: str | Path, sheets: Mapping[str, pd.DataFrame], metadata: Mapping[str, object] | None = None) -> Path:
        output = Path(path)
        output.parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            if metadata:
                pd.DataFrame([{"key": key, "value": value} for key, value in metadata.items()]).to_excel(
                    writer, sheet_name="Metadata", index=False
                )
            for name, frame in sheets.items():
                safe_name = self._sheet_name(name)
                frame.to_excel(writer, sheet_name=safe_name, index=False)
            workbook = writer.book
            for worksheet in workbook.worksheets:
                self._format_sheet(worksheet)
            if "Breadth History" in workbook.sheetnames:
                self._add_breadth_chart(workbook["Breadth History"])
        return output

    @staticmethod
    def _sheet_name(name: str) -> str:
        invalid = set('[]:*?/\\')
        cleaned = "".join("_" if char in invalid else char for char in name).strip()
        return cleaned[:31] or "Sheet"

    @staticmethod
    def _format_sheet(worksheet: object) -> None:
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_cells in worksheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column_cells[:200]]
            width = min(max(max((len(value) for value in values), default=0) + 2, 10), 40)
            worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        headers = {cell.value: cell.column for cell in worksheet[1]}
        score_column = headers.get("behavior_dna_score")
        if score_column and worksheet.max_row > 1:
            letter = get_column_letter(score_column)
            worksheet.conditional_formatting.add(
                f"{letter}2:{letter}{worksheet.max_row}",
                ColorScaleRule(start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=50, mid_color="FFEB84", end_type="num", end_value=100, end_color="63BE7B"),
            )

    @staticmethod
    def _add_breadth_chart(worksheet: object) -> None:
        headers = {cell.value: cell.column for cell in worksheet[1]}
        date_col = headers.get("date")
        breadth_col = headers.get("breadth_pct")
        if not date_col or not breadth_col or worksheet.max_row < 3:
            return
        chart = LineChart()
        chart.title = "BIST Market Breadth"
        chart.y_axis.title = "Advancing Share (%)"
        chart.x_axis.title = "Date"
        data = Reference(worksheet, min_col=breadth_col, min_row=1, max_row=worksheet.max_row)
        categories = Reference(worksheet, min_col=date_col, min_row=2, max_row=worksheet.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.height = 8
        chart.width = 16
        worksheet.add_chart(chart, "A4")
